from __future__ import annotations

import csv
import logging
import os
import shutil
import tempfile
from abc import ABC, abstractmethod
from contextlib import suppress
from dataclasses import dataclass
from datetime import UTC, datetime, time, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from avito_crm.config import Settings
from avito_crm.errors import ConfigurationError, SourceError
from avito_crm.google_api import google_api_call
from avito_crm.models import TERMINAL_STATUSES, ItemStatus, QueueItem, QueuePatch
from avito_crm.phone import canonical_avito_url

LOGGER = logging.getLogger(__name__)
_PLAN_PRIORITY = "__plan_priority"
_MOSCOW_OFFSET = "__moscow_offset"
_LEGACY_FALSE_CAPTCHA_ERROR = (
    "ручная проверка avito не завершена за отведённое время"
)


@dataclass(frozen=True, slots=True)
class QueueColumns:
    url: str
    status: str
    phone: str
    crm_lead_id: str
    funnel_stage: str
    crm_create_count: str
    repeat_crm_lead_id: str
    repeat_phone_attempts: str
    error: str
    attempts: str
    processed_at: str
    run_id: str
    next_retry_at: str

    @classmethod
    def from_settings(cls, settings: Settings) -> QueueColumns:
        return cls(
            url=settings.url_column,
            status=settings.status_column,
            phone=settings.phone_column,
            crm_lead_id=settings.crm_lead_column,
            funnel_stage=settings.funnel_stage_column,
            crm_create_count=settings.crm_create_count_column,
            repeat_crm_lead_id=settings.repeat_crm_lead_column,
            repeat_phone_attempts=settings.repeat_phone_attempts_column,
            error=settings.error_column,
            attempts=settings.attempts_column,
            processed_at=settings.processed_at_column,
            run_id=settings.run_id_column,
            next_retry_at=settings.next_retry_at_column,
        )

    @property
    def managed(self) -> tuple[str, ...]:
        return (
            self.status,
            self.phone,
            self.crm_lead_id,
            self.funnel_stage,
            self.crm_create_count,
            self.repeat_crm_lead_id,
            self.repeat_phone_attempts,
            self.error,
            self.attempts,
            self.processed_at,
            self.run_id,
            self.next_retry_at,
        )


class QueueSource(ABC):
    def __init__(
        self, columns: QueueColumns, max_attempts: int, repeat_phone_max_attempts: int = 3
    ) -> None:
        self.columns = columns
        self.max_attempts = max_attempts
        self.repeat_phone_max_attempts = repeat_phone_max_attempts

    def list_all(self) -> list[QueueItem]:
        raise NotImplementedError

    @abstractmethod
    def list_actionable(self, *, include_manual: bool = False) -> list[QueueItem]:
        raise NotImplementedError

    @abstractmethod
    def update(self, item: QueueItem, patch: QueuePatch) -> None:
        raise NotImplementedError

    def moscow_offset(self, item: QueueItem) -> int:
        """Return the row's Moscow offset; absent/invalid values safely map to zero."""
        return _safe_int(item.values.get(_MOSCOW_OFFSET))

    def _is_actionable(
        self,
        status: str,
        attempts: int,
        include_manual: bool,
        crm_create_count: int = 0,
        repeat_phone_attempts: int = 0,
        next_retry_at: str = "",
        error: str = "",
    ) -> bool:
        normalized = (status or "").strip().lower()
        # Extension versions before 1.0.6 could mistake an ordinary unopened
        # phone for a captcha and eventually store this exact timeout. Recover
        # only that known legacy marker; genuine manual_required rows must
        # continue waiting for the operator.
        if (
            normalized == ItemStatus.MANUAL_REQUIRED
            and _LEGACY_FALSE_CAPTCHA_ERROR in str(error or "").casefold()
        ):
            normalized = ItemStatus.RETRY_PHONE.value
        if crm_create_count >= 2 or crm_create_count not in {0, 1}:
            return False
        if crm_create_count == 1:
            return (
                normalized
                in {
                    ItemStatus.REPEAT_PENDING.value,
                    ItemStatus.REPEAT_RETRY_PHONE.value,
                    ItemStatus.REPEAT_RETRY_TECHNICAL.value,
                }
                and repeat_phone_attempts < getattr(self, "repeat_phone_max_attempts", 3)
                and (attempts < self.max_attempts or not str(next_retry_at or "").strip())
                and self._retry_is_due(normalized, next_retry_at)
            )
        if normalized in TERMINAL_STATUSES:
            return False
        if normalized == ItemStatus.MANUAL_REQUIRED and not include_manual:
            return False
        if not self._retry_is_due(normalized, next_retry_at):
            return False
        # Rows left by older versions have no explicit retry timestamp. Give
        # each such retry row exactly one final low-priority session, even when
        # its legacy counter used different semantics.
        if (
            normalized
            in {
                ItemStatus.RETRY_PHONE.value,
                ItemStatus.RETRY_TECHNICAL.value,
            }
            and not str(next_retry_at or "").strip()
        ):
            return True
        return attempts < self.max_attempts

    @staticmethod
    def _retry_is_due(status: str, next_retry_at: str) -> bool:
        retry_statuses = {
            ItemStatus.RETRY_PHONE.value,
            ItemStatus.RETRY_TECHNICAL.value,
            ItemStatus.REPEAT_RETRY_PHONE.value,
            ItemStatus.REPEAT_RETRY_TECHNICAL.value,
        }
        if status not in retry_statuses or not str(next_retry_at or "").strip():
            return True
        try:
            due = datetime.fromisoformat(str(next_retry_at).strip().replace("Z", "+00:00"))
            if due.tzinfo is None:
                due = due.replace(tzinfo=UTC)
        except ValueError:
            return True
        return datetime.now(UTC) >= due.astimezone(UTC)

    def _sorted_actionable(
        self, items: list[QueueItem], *, include_manual: bool
    ) -> list[QueueItem]:
        actionable = [
            item
            for item in items
            if self._is_actionable(
                item.status,
                item.attempts,
                include_manual,
                _safe_int(item.values.get(self.columns.crm_create_count)),
                _safe_int(item.values.get(self.columns.repeat_phone_attempts)),
                str(item.values.get(self.columns.next_retry_at, "") or ""),
                str(item.values.get(self.columns.error, "") or ""),
            )
        ]

        def key(item: QueueItem) -> tuple[int, int, int]:
            status = (item.status or "").strip().lower()
            if (
                status == ItemStatus.MANUAL_REQUIRED
                and _LEGACY_FALSE_CAPTCHA_ERROR
                in str(item.values.get(self.columns.error, "") or "").casefold()
            ):
                status = ItemStatus.RETRY_PHONE.value
            if status in {
                ItemStatus.RETRY_PHONE.value,
                ItemStatus.RETRY_TECHNICAL.value,
                ItemStatus.REPEAT_RETRY_PHONE.value,
                ItemStatus.REPEAT_RETRY_TECHNICAL.value,
            }:
                lane = 2
            elif status == ItemStatus.REPEAT_PENDING.value:
                lane = 1
            else:
                lane = 0
            priority = _safe_int(item.values.get(_PLAN_PRIORITY)) or 1_000_000
            return lane, priority, _safe_int(item.row_id)

        return sorted(actionable, key=key)

    def is_local_window_open(self, item: QueueItem, *, now: datetime | None = None) -> bool:
        return True

    def local_window_detail(self, item: QueueItem, *, now: datetime | None = None) -> str:
        return "ограничение местного времени для этого источника не задано"

    def next_local_window_at(
        self, items: list[QueueItem], *, now: datetime | None = None
    ) -> datetime | None:
        """Return the earliest safe opening for deferred rows, if one is known."""
        del items, now
        return None


class XlsxQueueSource(QueueSource):
    def __init__(
        self,
        path: Path,
        worksheet: str,
        columns: QueueColumns,
        max_attempts: int,
        backup_dir: Path,
        repeat_phone_max_attempts: int = 3,
    ) -> None:
        super().__init__(columns, max_attempts, repeat_phone_max_attempts)
        self.path = path.resolve()
        self.worksheet = worksheet
        self.backup_dir = backup_dir
        self._backup_done = False
        if not self.path.is_file():
            raise SourceError(f"Excel-файл не найден: {self.path}")

    def _load(self):
        try:
            workbook = load_workbook(self.path)
        except Exception as exc:
            raise SourceError(f"Не удалось открыть Excel-файл: {exc}") from exc
        if self.worksheet not in workbook.sheetnames:
            workbook.close()
            raise SourceError(
                f"Лист {self.worksheet!r} не найден; доступны: {', '.join(workbook.sheetnames)}"
            )
        return workbook, workbook[self.worksheet]

    def _headers(self, sheet) -> dict[str, int]:
        headers: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(1, column).value
            if value is not None and str(value).strip():
                headers[str(value).strip()] = column
        if self.columns.url not in headers:
            raise SourceError(f"В первой строке нет колонки {self.columns.url!r}")
        next_column = max(headers.values(), default=0) + 1
        for name in self.columns.managed:
            if name not in headers:
                sheet.cell(1, next_column, name)
                headers[name] = next_column
                next_column += 1
        return headers

    def _backup(self) -> None:
        if self._backup_done:
            return
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
        destination = self.backup_dir / f"{self.path.stem}.{stamp}{self.path.suffix}"
        shutil.copy2(self.path, destination)
        self._backup_done = True

    def _atomic_save(self, workbook) -> None:
        self._backup()
        fd, temp_name = tempfile.mkstemp(
            prefix=f".{self.path.stem}.", suffix=self.path.suffix, dir=self.path.parent
        )
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            workbook.save(temp_path)
            os.replace(temp_path, self.path)
        except Exception as exc:
            temp_path.unlink(missing_ok=True)
            raise SourceError(f"Не удалось сохранить Excel-файл: {exc}") from exc

    def list_all(self) -> list[QueueItem]:
        workbook, sheet = self._load()
        try:
            original_headers = {str(cell.value or "").strip() for cell in sheet[1]}
            headers = self._headers(sheet)
            headers_changed = any(name not in original_headers for name in self.columns.managed)
            items: list[QueueItem] = []
            for row in range(2, sheet.max_row + 1):
                url = str(sheet.cell(row, headers[self.columns.url]).value or "").strip()
                if not url:
                    continue
                status = str(sheet.cell(row, headers[self.columns.status]).value or "").strip()
                attempts = _safe_int(sheet.cell(row, headers[self.columns.attempts]).value)
                values = {name: sheet.cell(row, column).value for name, column in headers.items()}
                items.append(QueueItem(str(row), url, status, attempts, values))
            if headers_changed:
                self._atomic_save(workbook)
            return items
        finally:
            workbook.close()

    def list_actionable(self, *, include_manual: bool = False) -> list[QueueItem]:
        return self._sorted_actionable(self.list_all(), include_manual=include_manual)

    def update(self, item: QueueItem, patch: QueuePatch) -> None:
        workbook, sheet = self._load()
        try:
            headers = self._headers(sheet)
            row = _safe_int(item.row_id)
            if row < 2 or row > sheet.max_row:
                raise SourceError(f"Строка очереди больше не существует: {item.row_id}")
            current_url = str(sheet.cell(row, headers[self.columns.url]).value or "").strip()
            if current_url != item.url:
                raise SourceError(
                    f"Ссылка в строке {row} изменилась во время обработки; результат не записан"
                )
            values = _patch_values(self.columns, patch)
            for name, value in values.items():
                sheet.cell(row, headers[name], value)
            self._atomic_save(workbook)
        finally:
            workbook.close()


class CsvQueueSource(QueueSource):
    def __init__(
        self,
        path: Path,
        columns: QueueColumns,
        max_attempts: int,
        backup_dir: Path,
        repeat_phone_max_attempts: int = 3,
    ) -> None:
        super().__init__(columns, max_attempts, repeat_phone_max_attempts)
        self.path = path.resolve()
        self.backup_dir = backup_dir
        self._backup_done = False
        if not self.path.is_file():
            raise SourceError(f"CSV-файл не найден: {self.path}")

    def _read(self) -> tuple[list[str], list[dict[str, str]], bool]:
        try:
            with self.path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                headers = [str(value).strip() for value in (reader.fieldnames or [])]
                rows = [dict(row) for row in reader]
        except Exception as exc:
            raise SourceError(f"Не удалось прочитать CSV: {exc}") from exc
        if self.columns.url not in headers:
            raise SourceError(f"В CSV нет колонки {self.columns.url!r}")
        changed = False
        for name in self.columns.managed:
            if name not in headers:
                headers.append(name)
                changed = True
        return headers, rows, changed

    def _write(self, headers: list[str], rows: list[dict[str, Any]]) -> None:
        if not self._backup_done:
            self.backup_dir.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
            shutil.copy2(self.path, self.backup_dir / f"{self.path.stem}.{stamp}.csv")
            self._backup_done = True
        fd, temp_name = tempfile.mkstemp(prefix=f".{self.path.stem}.", dir=self.path.parent)
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            with temp_path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
                writer.writeheader()
                writer.writerows(rows)
            os.replace(temp_path, self.path)
        except Exception as exc:
            temp_path.unlink(missing_ok=True)
            raise SourceError(f"Не удалось сохранить CSV: {exc}") from exc

    def list_all(self) -> list[QueueItem]:
        headers, rows, headers_changed = self._read()
        if headers_changed:
            self._write(headers, rows)
        items = []
        for index, row in enumerate(rows, start=2):
            url = str(row.get(self.columns.url, "") or "").strip()
            if not url:
                continue
            status = str(row.get(self.columns.status, "") or "").strip()
            attempts = _safe_int(row.get(self.columns.attempts))
            items.append(QueueItem(str(index), url, status, attempts, row))
        return items

    def list_actionable(self, *, include_manual: bool = False) -> list[QueueItem]:
        return self._sorted_actionable(self.list_all(), include_manual=include_manual)

    def update(self, item: QueueItem, patch: QueuePatch) -> None:
        headers, rows, _headers_changed = self._read()
        index = _safe_int(item.row_id) - 2
        if index < 0 or index >= len(rows):
            raise SourceError(f"Строка очереди больше не существует: {item.row_id}")
        if str(rows[index].get(self.columns.url, "") or "").strip() != item.url:
            raise SourceError(f"Ссылка в строке {item.row_id} изменилась; результат не записан")
        rows[index].update(_patch_values(self.columns, patch))
        self._write(headers, rows)


class GoogleSheetsQueueSource(QueueSource):
    def __init__(
        self,
        credentials_file: Path,
        spreadsheet_id: str,
        worksheet: str,
        columns: QueueColumns,
        max_attempts: int,
        repeat_phone_max_attempts: int = 3,
        *,
        plan_worksheet: str = "План загрузки",
        timezone_guard_enabled: bool = False,
        local_call_start: time = time(10, 0),
        local_lead_cutoff: time = time(19, 45),
    ) -> None:
        super().__init__(columns, max_attempts, repeat_phone_max_attempts)
        if not credentials_file.is_file():
            raise ConfigurationError(
                f"Файл сервисного аккаунта Google не найден: {credentials_file}"
            )
        if not spreadsheet_id:
            raise ConfigurationError("Не задан GOOGLE_SPREADSHEET_ID")
        try:
            import gspread

            client = gspread.service_account(filename=str(credentials_file))
            spreadsheet = google_api_call(
                lambda: client.open_by_key(spreadsheet_id),
                label="Открытие Google Sheet очереди",
            )
            self.spreadsheet = spreadsheet
            self.sheet = google_api_call(
                lambda: spreadsheet.worksheet(worksheet),
                label=f"Открытие листа очереди {worksheet!r}",
            )
            self._headers_cache: list[str] | None = None
            self._rows_cache: list[list[str]] | None = None
            self.plan_worksheet = plan_worksheet
            self.timezone_guard_enabled = timezone_guard_enabled
            self.local_call_start = local_call_start
            self.local_lead_cutoff = local_lead_cutoff
            self._plan_cache: dict[str, tuple[int, int]] | None = None
        except Exception as exc:
            raise SourceError(f"Не удалось открыть Google Sheet: {exc}") from exc

    def _read(self) -> tuple[list[str], list[list[str]]]:
        try:
            values = google_api_call(
                self.sheet.get_all_values,
                label="Чтение Google Sheet очереди",
            )
        except Exception as exc:
            raise SourceError(f"Не удалось прочитать Google Sheet: {exc}") from exc
        if not values:
            headers = [self.columns.url, *self.columns.managed]
            try:
                google_api_call(
                    lambda: self.sheet.update([headers], "1:1", value_input_option="RAW"),
                    label="Подготовка заголовков Google Sheet",
                )
            except Exception as exc:
                raise SourceError(f"Не удалось подготовить пустой Google Sheet: {exc}") from exc
            self._format_empty_sheet(headers)
            self._headers_cache = headers
            self._rows_cache = []
            return headers, []
        headers = [str(value).strip() for value in values[0]]
        if self.columns.url not in headers:
            raise SourceError(f"В Google Sheet нет колонки {self.columns.url!r}")
        changed = False
        for name in self.columns.managed:
            if name not in headers:
                headers.append(name)
                changed = True
        if changed:
            try:
                google_api_call(
                    lambda: self.sheet.update([headers], "1:1", value_input_option="RAW"),
                    label="Добавление служебных колонок Google Sheet",
                )
            except Exception as exc:
                raise SourceError(f"Не удалось добавить служебные колонки: {exc}") from exc
        width = len(headers)
        rows = [row + [""] * (width - len(row)) for row in values[1:]]
        self._headers_cache = headers
        self._rows_cache = rows
        return headers, rows

    def _format_empty_sheet(self, headers: list[str]) -> None:
        """Apply a restrained native layout only when the tab was fully empty."""
        with suppress(Exception):
            self.sheet.freeze(rows=1)
            self.sheet.format(
                "1:1",
                {
                    "backgroundColor": {"red": 0.95, "green": 0.96, "blue": 0.97},
                    "textFormat": {
                        "bold": True,
                        "foregroundColor": {"red": 0.1, "green": 0.13, "blue": 0.19},
                    },
                    "verticalAlignment": "MIDDLE",
                    "wrapStrategy": "CLIP",
                },
            )
            self.sheet.set_basic_filter()

        spreadsheet = getattr(self, "spreadsheet", None)
        sheet_id = getattr(self.sheet, "id", None)
        if spreadsheet is None or sheet_id is None:
            return
        with suppress(Exception):
            spreadsheet.batch_update(
                {
                    "requests": [
                        {
                            "updateDimensionProperties": {
                                "range": {
                                    "sheetId": sheet_id,
                                    "dimension": "COLUMNS",
                                    "startIndex": 0,
                                    "endIndex": 1,
                                },
                                "properties": {"pixelSize": 520},
                                "fields": "pixelSize",
                            }
                        },
                        {
                            "updateDimensionProperties": {
                                "range": {
                                    "sheetId": sheet_id,
                                    "dimension": "COLUMNS",
                                    "startIndex": 1,
                                    "endIndex": len(headers),
                                },
                                "properties": {"pixelSize": 150},
                                "fields": "pixelSize",
                            }
                        },
                    ]
                }
            )

    def list_all(self) -> list[QueueItem]:
        headers, rows = self._read()
        index = {name: position for position, name in enumerate(headers)}
        plan = self._read_plan() if getattr(self, "timezone_guard_enabled", False) else {}
        items = []
        for row_number, row in enumerate(rows, start=2):
            url = row[index[self.columns.url]].strip()
            if not url:
                continue
            status = row[index[self.columns.status]].strip()
            attempts = _safe_int(row[index[self.columns.attempts]])
            values = {name: row[position] for name, position in index.items()}
            plan_entry = plan.get(_canonical_or_raw(url))
            if plan_entry is not None:
                values[_PLAN_PRIORITY], values[_MOSCOW_OFFSET] = plan_entry
            items.append(QueueItem(str(row_number), url, status, attempts, values))
        return items

    def list_actionable(self, *, include_manual: bool = False) -> list[QueueItem]:
        return self._sorted_actionable(self.list_all(), include_manual=include_manual)

    def _read_plan(self) -> dict[str, tuple[int, int]]:
        if self._plan_cache is not None:
            return self._plan_cache
        try:
            sheet = google_api_call(
                lambda: self.spreadsheet.worksheet(self.plan_worksheet),
                label=f"Открытие листа планирования {self.plan_worksheet!r}",
            )
            values = google_api_call(
                sheet.get_all_values,
                label="Чтение приоритетов и часовых поясов",
            )
        except Exception as exc:
            raise SourceError(f"Не удалось прочитать план часовых поясов: {exc}") from exc
        required = {"Приоритет", "Ссылка", "Δ к МСК"}
        header_index = next(
            (
                index
                for index, row in enumerate(values[:10])
                if required.issubset({str(value).strip() for value in row})
            ),
            None,
        )
        if header_index is None:
            raise SourceError(
                f"В листе {self.plan_worksheet!r} не найдены колонки: "
                + ", ".join(sorted(required))
            )
        headers = [str(value).strip() for value in values[header_index]]
        columns = {name: headers.index(name) for name in required}
        result: dict[str, tuple[int, int]] = {}
        for row in values[header_index + 1 :]:
            padded = row + [""] * (len(headers) - len(row))
            url = str(padded[columns["Ссылка"]]).strip()
            if not url:
                continue
            try:
                priority = int(float(str(padded[columns["Приоритет"]]).strip()))
                offset = int(float(str(padded[columns["Δ к МСК"]]).strip()))
            except ValueError:
                continue
            result[_canonical_or_raw(url)] = (priority, offset)
        self._plan_cache = result
        return result

    def is_local_window_open(self, item: QueueItem, *, now: datetime | None = None) -> bool:
        if not self.timezone_guard_enabled:
            return True
        raw_offset = item.values.get(_MOSCOW_OFFSET)
        if raw_offset is None:
            return False
        moscow = ZoneInfo("Europe/Moscow")
        current = now or datetime.now(moscow)
        if current.tzinfo is None:
            current = current.replace(tzinfo=moscow)
        local = current.astimezone(moscow) + timedelta(hours=_safe_int(raw_offset))
        local_clock = local.timetz().replace(tzinfo=None)
        return self.local_call_start <= local_clock <= self.local_lead_cutoff

    def local_window_detail(self, item: QueueItem, *, now: datetime | None = None) -> str:
        raw_offset = item.values.get(_MOSCOW_OFFSET)
        if raw_offset is None:
            return "URL отсутствует в листе планирования; безопасное время неизвестно"
        moscow = ZoneInfo("Europe/Moscow")
        current = now or datetime.now(moscow)
        if current.tzinfo is None:
            current = current.replace(tzinfo=moscow)
        local = current.astimezone(moscow) + timedelta(hours=_safe_int(raw_offset))
        return (
            f"местное время {local:%H:%M}, разрешено "
            f"{self.local_call_start:%H:%M}–{self.local_lead_cutoff:%H:%M}"
        )

    def next_local_window_at(
        self, items: list[QueueItem], *, now: datetime | None = None
    ) -> datetime | None:
        if not self.timezone_guard_enabled:
            return now or datetime.now(UTC)
        moscow = ZoneInfo("Europe/Moscow")
        current = now or datetime.now(moscow)
        if current.tzinfo is None:
            current = current.replace(tzinfo=moscow)
        current_moscow = current.astimezone(moscow)
        candidates: list[datetime] = []
        for item in items:
            raw_offset = item.values.get(_MOSCOW_OFFSET)
            if raw_offset is None:
                continue
            offset = timedelta(hours=_safe_int(raw_offset))
            local_now = current_moscow + offset
            local_date = local_now.date()
            if local_now.timetz().replace(tzinfo=None) >= self.local_call_start:
                local_date += timedelta(days=1)
            local_open = datetime.combine(local_date, self.local_call_start, tzinfo=moscow)
            candidates.append(local_open - offset)
        return min(candidates) if candidates else None

    def update(self, item: QueueItem, patch: QueuePatch) -> None:
        from gspread.utils import rowcol_to_a1

        headers = getattr(self, "_headers_cache", None)
        rows = getattr(self, "_rows_cache", None)
        if headers is None or rows is None:
            headers, rows = self._read()
        index = {name: position for position, name in enumerate(headers)}
        row_number = _safe_int(item.row_id)
        values_index = row_number - 2
        if values_index < 0 or values_index >= len(rows):
            raise SourceError(f"Строка Google Sheet больше не существует: {item.row_id}")
        url_column = index[self.columns.url]
        current_url = rows[values_index][url_column].strip()
        if hasattr(self.sheet, "acell"):
            try:
                url_cell = rowcol_to_a1(row_number, url_column + 1)
                current_url = str(
                    google_api_call(
                        lambda: self.sheet.acell(url_cell).value,
                        label=f"Проверка строки Google Sheet {row_number}",
                    )
                    or ""
                ).strip()
            except Exception as exc:
                raise SourceError(f"Не удалось проверить строку {item.row_id}: {exc}") from exc
        if current_url != item.url:
            raise SourceError(f"Ссылка в строке {item.row_id} изменилась; результат не записан")
        patch_values = _patch_values(self.columns, patch)
        cells = []
        for name, value in patch_values.items():
            cells.append((rowcol_to_a1(row_number, index[name] + 1), value))

        def update_cells() -> Any:
            # gspread qualifies every relative A1 range with the worksheet title and
            # mutates the supplied dictionaries in place.  google_api_call may invoke
            # this operation again after a transient 429/5xx response, so each attempt
            # must receive a fresh payload or a range such as B401 becomes
            # 'Sheet'!'Sheet'!B401 on the retry.
            payload = [{"range": range_name, "values": [[value]]} for range_name, value in cells]
            return self.sheet.batch_update(payload, value_input_option="RAW")

        try:
            google_api_call(
                update_cells,
                label=f"Обновление строки Google Sheet {row_number}",
            )
        except Exception as exc:
            raise SourceError(f"Не удалось обновить Google Sheet: {exc}") from exc
        for name, value in patch_values.items():
            rows[values_index][index[name]] = str(value or "")


def build_queue_source(
    settings: Settings,
    source_type: str,
    file_path: Path | None,
    worksheet: str | None,
) -> QueueSource:
    columns = QueueColumns.from_settings(settings)
    source_type = source_type.lower()
    sheet_name = worksheet or settings.google_worksheet
    backup_dir = settings.data_dir / "backups"
    if source_type == "xlsx":
        if file_path is None:
            raise ConfigurationError("Для source=xlsx нужен параметр --file")
        return XlsxQueueSource(
            file_path,
            sheet_name,
            columns,
            settings.max_attempts,
            backup_dir,
            settings.repeat_phone_max_attempts,
        )
    if source_type == "csv":
        if file_path is None:
            raise ConfigurationError("Для source=csv нужен параметр --file")
        return CsvQueueSource(
            file_path,
            columns,
            settings.max_attempts,
            backup_dir,
            settings.repeat_phone_max_attempts,
        )
    if source_type == "google":
        if settings.google_credentials_file is None:
            raise ConfigurationError("Не задан GOOGLE_CREDENTIALS_FILE")
        source = GoogleSheetsQueueSource(
            settings.google_credentials_file,
            settings.google_spreadsheet_id,
            sheet_name,
            columns,
            settings.max_attempts,
            settings.repeat_phone_max_attempts,
            plan_worksheet=settings.google_plan_worksheet,
            timezone_guard_enabled=settings.timezone_guard_enabled,
            local_call_start=settings.local_call_start,
            local_lead_cutoff=settings.local_lead_cutoff,
        )
        return source
    raise ConfigurationError("--source: допустимо google, xlsx или csv")


def _safe_int(value: Any) -> int:
    try:
        return int(float(str(value or "0").strip()))
    except (TypeError, ValueError):
        return 0


def _patch_values(columns: QueueColumns, patch: QueuePatch) -> dict[str, Any]:
    values = {
        columns.status: patch.status,
        columns.phone: patch.phone,
        columns.crm_lead_id: patch.crm_lead_id,
        columns.error: (patch.error or "")[:500],
        columns.attempts: patch.attempts,
        columns.processed_at: patch.processed_at,
        columns.run_id: patch.run_id,
    }
    optional = {
        columns.funnel_stage: patch.funnel_stage,
        columns.crm_create_count: patch.crm_create_count,
        columns.repeat_crm_lead_id: patch.repeat_crm_lead_id,
        columns.repeat_phone_attempts: patch.repeat_phone_attempts,
        columns.next_retry_at: patch.next_retry_at,
    }
    values.update({name: value for name, value in optional.items() if value is not None})
    return values


def _canonical_or_raw(url: str) -> str:
    try:
        return canonical_avito_url(url)
    except Exception:
        return str(url or "").strip()
