from __future__ import annotations

import csv
import hashlib
import json
import random
import re
import time
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from avito_crm.chrome_extension import ChromeExtensionBrowser
from avito_crm.config import Settings
from avito_crm.errors import (
    BrowserOperationError,
    ConfigurationError,
    InactiveListingError,
    InvalidListingError,
    ManualActionRequired,
    OperatorStopRequested,
    PageNotReadyError,
    PhoneButtonUnavailableError,
    PhoneNotFoundError,
    SourceError,
)
from avito_crm.ocr import PhoneOcr
from avito_crm.phone import canonical_avito_url

MAX_BATCH_INPUT_BYTES = 10 * 1024 * 1024
MAX_BATCH_REPORT_BYTES = 20 * 1024 * 1024
MAX_BATCH_URLS = 500
URL_COLUMN_NAMES = {"url", "link", "ссылка", "объявление", "ссылка avito"}
SAFE_URL_HASH_RE = re.compile(r"[0-9a-f]{16}")


@dataclass(frozen=True, slots=True)
class BatchUrl:
    ordinal: int
    input_row: int
    url: str


@dataclass(slots=True)
class ExtensionBatchSummary:
    requested: int
    report_path: Path
    inspected: int = 0
    succeeded: int = 0
    stopped_reason: str = ""
    statuses: Counter[str] = field(default_factory=Counter)
    sources: Counter[str] = field(default_factory=Counter)
    elapsed_seconds: float = 0.0

    @property
    def failed(self) -> int:
        return self.inspected - self.succeeded


def load_batch_urls(path: Path, *, limit: int, sheet: str | None = None) -> list[BatchUrl]:
    resolved = path.expanduser().resolve()
    if not resolved.is_file():
        raise ConfigurationError(f"Не найден файл ссылок: {resolved}")
    if resolved.stat().st_size > MAX_BATCH_INPUT_BYTES:
        raise ConfigurationError("Файл ссылок превышает безопасный предел 10 МБ")
    if not 1 <= limit <= MAX_BATCH_URLS:
        raise ConfigurationError(f"--limit должен быть от 1 до {MAX_BATCH_URLS}")

    suffix = resolved.suffix.casefold()
    if suffix == ".txt":
        rows = _read_text_rows(resolved)
    elif suffix == ".csv":
        rows = _read_csv_rows(resolved)
    elif suffix == ".xlsx":
        rows = _read_xlsx_rows(resolved, sheet=sheet)
    else:
        raise ConfigurationError("Для batch-теста поддерживаются только .txt, .csv и .xlsx")

    selected: list[BatchUrl] = []
    seen: set[str] = set()
    invalid: list[int] = []
    for input_row, raw_url in rows:
        try:
            url = canonical_avito_url(raw_url)
        except InvalidListingError:
            invalid.append(input_row)
            continue
        if url in seen:
            continue
        seen.add(url)
        selected.append(BatchUrl(len(selected) + 1, input_row, url))
        if len(selected) >= limit:
            break

    if invalid:
        preview = ", ".join(str(value) for value in invalid[:5])
        raise ConfigurationError("Файл содержит некорректные ссылки Avito в строках: " + preview)
    if len(selected) < limit:
        raise ConfigurationError(
            f"В файле только {len(selected)} уникальных корректных ссылок; запрошено {limit}"
        )
    return selected


def load_google_batch_urls(
    settings: Settings,
    *,
    limit: int,
    sheet: str | None = None,
    statuses: tuple[str, ...] = ("retry_phone",),
    excluded_url_hashes: set[str] | None = None,
) -> list[BatchUrl]:
    """Read test URLs from Google without constructing a writable queue source."""
    if not 1 <= limit <= MAX_BATCH_URLS:
        raise ConfigurationError(f"--limit должен быть от 1 до {MAX_BATCH_URLS}")
    if settings.google_credentials_file is None:
        raise ConfigurationError("Не задан GOOGLE_CREDENTIALS_FILE")
    if not settings.google_credentials_file.is_file():
        raise ConfigurationError(
            f"Файл сервисного аккаунта Google не найден: {settings.google_credentials_file}"
        )
    if not settings.google_spreadsheet_id:
        raise ConfigurationError("Не задан GOOGLE_SPREADSHEET_ID")
    normalized_statuses = {value.strip().casefold() for value in statuses if value.strip()}
    if not normalized_statuses:
        raise ConfigurationError("Для Google batch-теста нужен хотя бы один --status")

    try:
        import gspread

        client = gspread.service_account(filename=str(settings.google_credentials_file))
        spreadsheet = client.open_by_key(settings.google_spreadsheet_id)
        worksheet = spreadsheet.worksheet(sheet or settings.google_worksheet)
        values = worksheet.get_all_values()
    except Exception as exc:
        raise SourceError(f"Не удалось прочитать Google Sheet для batch-теста: {exc}") from exc

    return _select_google_batch_rows(
        values,
        url_column=settings.url_column,
        status_column=settings.status_column,
        statuses=normalized_statuses,
        limit=limit,
        excluded_url_hashes=excluded_url_hashes,
    )


def load_tested_url_hashes(output_dir: Path) -> set[str]:
    """Read safe URL hashes from completed or partial local batch reports."""
    hashes: set[str] = set()
    for report in sorted(output_dir.glob("extension-batch-*.jsonl")):
        if not report.is_file():
            continue
        if report.stat().st_size > MAX_BATCH_REPORT_BYTES:
            raise SourceError(f"Batch-отчёт превышает безопасный предел 20 МБ: {report.name}")
        try:
            with report.open("r", encoding="utf-8") as stream:
                for line_number, line in enumerate(stream, 1):
                    if not line.strip():
                        continue
                    payload = json.loads(line)
                    if not isinstance(payload, dict):
                        raise SourceError(
                            f"Batch-отчёт содержит запись неверного формата: "
                            f"{report.name}, строка {line_number}"
                        )
                    if payload.get("record") != "item":
                        continue
                    url_hash = str(payload.get("url_sha256", "")).strip().casefold()
                    if not SAFE_URL_HASH_RE.fullmatch(url_hash):
                        raise SourceError(
                            f"Batch-отчёт не содержит безопасный URL-хэш: "
                            f"{report.name}, строка {line_number}"
                        )
                    hashes.add(url_hash)
        except (OSError, json.JSONDecodeError) as exc:
            raise SourceError(f"Не удалось безопасно прочитать batch-отчёт {report.name}") from exc
    return hashes


def run_extension_batch(
    settings: Settings,
    urls: list[BatchUrl],
    *,
    max_clicks: int,
    circuit_breaker: int,
    ocr: PhoneOcr | None = None,
    browser_factory: Callable[[Settings, PhoneOcr], ChromeExtensionBrowser] = (
        ChromeExtensionBrowser
    ),
    sleep: Callable[[float], None] = time.sleep,
    uniform: Callable[[float, float], float] = random.uniform,
    monotonic: Callable[[], float] = time.monotonic,
) -> ExtensionBatchSummary:
    if max_clicks not in {1, 2}:
        raise ConfigurationError("--max-clicks должен быть равен 1 или 2")
    if not 3 <= circuit_breaker <= 50:
        raise ConfigurationError("--circuit-breaker должен быть от 3 до 50")
    if not urls:
        raise ConfigurationError("Список ссылок для batch-теста пуст")

    report_path = _new_report_path(settings.output_dir)
    summary = ExtensionBatchSummary(requested=len(urls), report_path=report_path)
    phone_ocr = ocr or PhoneOcr(settings.tesseract_cmd, settings.ocr_min_agreement)
    phone_ocr.check_available()
    stop_file = settings.data_dir / "STOP"
    stop_file.unlink(missing_ok=True)

    started_at = monotonic()
    next_long_break_at = started_at + uniform(
        settings.avito_long_break_interval_min,
        settings.avito_long_break_interval_max,
    )
    consecutive_ocr = 0
    consecutive_technical = 0
    _append_jsonl(
        report_path,
        {
            "record": "start",
            "schema": 1,
            "started_at": _utc_now(),
            "requested": len(urls),
            "max_clicks": max_clicks,
            "circuit_breaker": circuit_breaker,
            "crm_writes": False,
            "queue_writes": False,
        },
    )

    try:
        with browser_factory(settings, phone_ocr) as browser:
            for index, item in enumerate(urls):
                if stop_file.exists():
                    summary.stopped_reason = "Остановка запрошена оператором"
                    break

                item_started = monotonic()
                status = "technical"
                source = ""
                try:
                    result = browser.reveal_phone(
                        item.url,
                        f"batch-{item.ordinal}",
                        max_clicks=max_clicks,
                    )
                except InactiveListingError:
                    status = "inactive"
                    consecutive_ocr = 0
                    consecutive_technical = 0
                except PhoneButtonUnavailableError:
                    status = "button_missing"
                    consecutive_ocr = 0
                    consecutive_technical = 0
                except PhoneNotFoundError:
                    status = "ocr_failed"
                    consecutive_ocr += 1
                    consecutive_technical = 0
                except PageNotReadyError:
                    status = "page_not_ready"
                    consecutive_technical += 1
                    consecutive_ocr = 0
                except InvalidListingError:
                    status = "invalid_listing"
                    consecutive_ocr = 0
                    consecutive_technical = 0
                except ManualActionRequired:
                    status = "manual_required"
                    summary.stopped_reason = "Chrome требует ручного действия оператора"
                except OperatorStopRequested:
                    status = "stopped"
                    summary.stopped_reason = "Остановка запрошена оператором"
                except BrowserOperationError:
                    status = "browser_error"
                    consecutive_technical += 1
                    consecutive_ocr = 0
                else:
                    status = "success"
                    source = result.source
                    summary.succeeded += 1
                    summary.sources[source] += 1
                    consecutive_ocr = 0
                    consecutive_technical = 0

                summary.inspected += 1
                summary.statuses[status] += 1
                _append_jsonl(
                    report_path,
                    {
                        "record": "item",
                        "ordinal": item.ordinal,
                        "input_row": item.input_row,
                        "url_sha256": hashlib.sha256(item.url.encode("utf-8")).hexdigest()[:16],
                        "status": status,
                        "source": source,
                        "elapsed_seconds": round(monotonic() - item_started, 3),
                    },
                )

                print(
                    f"[{summary.inspected}/{summary.requested}] {status}; "
                    f"успешно {summary.succeeded}; ошибок {summary.failed}"
                )
                if summary.stopped_reason:
                    break
                if consecutive_ocr >= circuit_breaker:
                    summary.stopped_reason = (
                        f"Circuit breaker: {consecutive_ocr} последовательных OCR-ошибок"
                    )
                    break
                if consecutive_technical >= circuit_breaker:
                    summary.stopped_reason = (
                        f"Circuit breaker: {consecutive_technical} последовательных ошибок Chrome"
                    )
                    break
                if index + 1 >= len(urls):
                    continue

                now = monotonic()
                if now >= next_long_break_at:
                    pause = uniform(
                        settings.avito_long_break_duration_min,
                        settings.avito_long_break_duration_max,
                    )
                    print(f"Плановая пауза batch-теста: {round(pause)} сек.")
                    sleep(pause)
                    next_long_break_at = monotonic() + uniform(
                        settings.avito_long_break_interval_min,
                        settings.avito_long_break_interval_max,
                    )
                else:
                    sleep(uniform(settings.avito_min_delay, settings.avito_max_delay))
    finally:
        summary.elapsed_seconds = max(0.0, monotonic() - started_at)
        _append_jsonl(
            report_path,
            {
                "record": "summary",
                "finished_at": _utc_now(),
                "requested": summary.requested,
                "inspected": summary.inspected,
                "succeeded": summary.succeeded,
                "failed": summary.failed,
                "success_rate": round(
                    summary.succeeded / summary.inspected if summary.inspected else 0.0,
                    4,
                ),
                "statuses": dict(summary.statuses),
                "sources": dict(summary.sources),
                "stopped_reason": summary.stopped_reason,
                "elapsed_seconds": round(summary.elapsed_seconds, 3),
                "crm_writes": False,
                "queue_writes": False,
            },
        )
    return summary


def _read_text_rows(path: Path) -> list[tuple[int, str]]:
    rows: list[tuple[int, str]] = []
    with path.open("r", encoding="utf-8-sig") as stream:
        for number, line in enumerate(stream, 1):
            value = line.strip()
            if value and not value.startswith("#"):
                rows.append((number, value))
    return rows


def _read_csv_rows(path: Path) -> list[tuple[int, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        table = list(csv.reader(stream))
    return _rows_from_table(table)


def _read_xlsx_rows(path: Path, *, sheet: str | None) -> list[tuple[int, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet:
            if sheet not in workbook.sheetnames:
                raise ConfigurationError(f"В Excel нет листа: {sheet}")
            worksheet = workbook[sheet]
        else:
            worksheet = workbook[workbook.sheetnames[0]]
        table = [[cell for cell in row] for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _rows_from_table(table)


def _rows_from_table(table: list[list[object]]) -> list[tuple[int, str]]:
    first_nonempty = next(
        (
            index
            for index, row in enumerate(table)
            if any(str(value or "").strip() for value in row)
        ),
        None,
    )
    if first_nonempty is None:
        return []
    header = [str(value or "").strip().casefold() for value in table[first_nonempty]]
    url_column = next(
        (index for index, value in enumerate(header) if value in URL_COLUMN_NAMES),
        0,
    )
    header_is_url = header[url_column].startswith(("http://", "https://"))
    start = first_nonempty if header_is_url else first_nonempty + 1
    rows: list[tuple[int, str]] = []
    for index in range(start, len(table)):
        row = table[index]
        value = str(row[url_column] or "").strip() if url_column < len(row) else ""
        if value:
            rows.append((index + 1, value))
    return rows


def _select_google_batch_rows(
    values: list[list[str]],
    *,
    url_column: str,
    status_column: str,
    statuses: set[str],
    limit: int,
    excluded_url_hashes: set[str] | None = None,
) -> list[BatchUrl]:
    if not values:
        raise SourceError("Google Sheet очереди пуст")
    headers = [str(value).strip() for value in values[0]]
    try:
        url_index = headers.index(url_column)
        status_index = headers.index(status_column)
    except ValueError as exc:
        raise SourceError(
            f"В Google Sheet нужны колонки {url_column!r} и {status_column!r}"
        ) from exc

    selected: list[BatchUrl] = []
    seen: set[str] = set()
    excluded = excluded_url_hashes or set()
    for row_number, row in enumerate(values[1:], 2):
        padded = row + [""] * max(0, len(headers) - len(row))
        if padded[status_index].strip().casefold() not in statuses:
            continue
        try:
            url = canonical_avito_url(padded[url_index])
        except InvalidListingError:
            continue
        url_hash = hashlib.sha256(url.encode("utf-8")).hexdigest()[:16]
        if url_hash in excluded:
            continue
        if url in seen:
            continue
        seen.add(url)
        selected.append(BatchUrl(len(selected) + 1, row_number, url))
        if len(selected) >= limit:
            break
    if len(selected) < limit:
        status_text = ", ".join(sorted(statuses))
        raise SourceError(
            f"В Google Sheet только {len(selected)} уникальных ссылок со статусами "
            f"{status_text}; запрошено {limit}"
        )
    return selected


def _new_report_path(output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%S%fZ")
    return output_dir / f"extension-batch-{stamp}.jsonl"


def _append_jsonl(path: Path, payload: dict[str, object]) -> None:
    with path.open("a", encoding="utf-8", newline="\n") as stream:
        stream.write(json.dumps(payload, ensure_ascii=False, sort_keys=True) + "\n")


def _utc_now() -> str:
    return datetime.now(UTC).isoformat(timespec="seconds")
