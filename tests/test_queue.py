from dataclasses import replace

import pytest
from openpyxl import Workbook, load_workbook

import avito_crm.queue as queue_module
from avito_crm.models import ItemStatus, QueuePatch
from avito_crm.queue import QueueColumns, XlsxQueueSource, build_queue_source


def test_xlsx_queue_adds_columns_updates_atomically_and_backs_up(tmp_path, settings):
    path = tmp_path / "queue.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"
    sheet.append(["Ссылка", "Комментарий"])
    sheet.append(["https://www.avito.ru/moskva/item_123456789", "keep me"])
    workbook.save(path)
    workbook.close()

    columns = QueueColumns.from_settings(settings)
    source = XlsxQueueSource(path, "Лист1", columns, 3, tmp_path / "backups")
    items = source.list_actionable()

    assert len(items) == 1
    source.update(
        items[0],
        QueuePatch(
            status=ItemStatus.CAPTURED,
            attempts=1,
            phone="+79991234567",
            processed_at="2026-07-18T00:00:00+00:00",
            run_id="run-1",
        ),
    )

    saved = load_workbook(path)
    saved_sheet = saved["Лист1"]
    headers = {cell.value: cell.column for cell in saved_sheet[1]}
    assert saved_sheet.cell(2, headers["Комментарий"]).value == "keep me"
    assert saved_sheet.cell(2, headers[columns.status]).value == ItemStatus.CAPTURED
    assert saved_sheet.cell(2, headers[columns.phone]).value == "+79991234567"
    saved.close()
    assert len(list((tmp_path / "backups").glob("*.xlsx"))) == 1


@pytest.mark.parametrize(
    "status",
    [
        ItemStatus.DONE,
        ItemStatus.DUPLICATE,
        ItemStatus.INACTIVE,
        ItemStatus.UNAVAILABLE,
        ItemStatus.NO_PHONE,
        ItemStatus.INVALID,
    ],
)
def test_terminal_xlsx_rows_are_not_actionable(tmp_path, settings, status):
    path = tmp_path / "queue.xlsx"
    columns = QueueColumns.from_settings(settings)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"
    sheet.append([columns.url, *columns.managed])
    sheet.append(["https://www.avito.ru/moskva/item_123456789", status, "+79991234567"])
    workbook.save(path)
    workbook.close()

    source = XlsxQueueSource(path, "Лист1", columns, 3, tmp_path / "backups")

    assert source.list_actionable() == []


def test_new_rows_are_processed_before_legacy_phone_errors(tmp_path, settings):
    path = tmp_path / "queue.xlsx"
    columns = QueueColumns.from_settings(settings)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"
    sheet.append([columns.url, *columns.managed])
    sheet.append(
        [
            "https://www.avito.ru/moskva/item_123456789",
            ItemStatus.RETRY_PHONE,
            "",
        ]
    )
    sheet.append(["https://www.avito.ru/moskva/item_123456790"])
    workbook.save(path)
    workbook.close()

    source = XlsxQueueSource(path, "Лист1", columns, 2, tmp_path / "backups")

    assert [item.row_id for item in source.list_actionable()] == ["3", "2"]


def test_legacy_false_captcha_timeout_returns_to_retry_lane(tmp_path, settings):
    path = tmp_path / "queue.xlsx"
    columns = QueueColumns.from_settings(settings)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"
    sheet.append([columns.url, *columns.managed])
    sheet.append(
        [
            "https://www.avito.ru/moskva/item_123456789",
            ItemStatus.MANUAL_REQUIRED,
            "",
            "",
            "",
            "",
            "",
            "",
            "ручная проверка Avito не завершена за отведённое время",
            2,
        ]
    )
    sheet.append(["https://www.avito.ru/moskva/item_123456790"])
    workbook.save(path)
    workbook.close()

    source = XlsxQueueSource(path, "Лист1", columns, 2, tmp_path / "backups")

    # The genuine new row stays first; the known false-captcha timeout gets
    # exactly the same low-priority recovery treatment as a phone retry.
    assert [item.row_id for item in source.list_actionable()] == ["3", "2"]


def test_genuine_manual_required_row_still_needs_explicit_requeue(tmp_path, settings):
    path = tmp_path / "queue.xlsx"
    columns = QueueColumns.from_settings(settings)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"
    sheet.append([columns.url, *columns.managed])
    sheet.append(
        [
            "https://www.avito.ru/moskva/item_123456789",
            ItemStatus.MANUAL_REQUIRED,
            "",
            "",
            "",
            "",
            "",
            "",
            "На странице обнаружена капча Avito",
        ]
    )
    workbook.save(path)
    workbook.close()

    source = XlsxQueueSource(path, "Лист1", columns, 2, tmp_path / "backups")

    assert source.list_actionable() == []
    assert [item.row_id for item in source.list_actionable(include_manual=True)] == ["2"]


def test_google_source_receives_plan_and_local_window_settings(tmp_path, settings, monkeypatch):
    captured = {}
    sentinel = object()

    def fake_google_source(*args, **kwargs):
        captured["args"] = args
        captured["kwargs"] = kwargs
        return sentinel

    monkeypatch.setattr(queue_module, "GoogleSheetsQueueSource", fake_google_source)
    configured = replace(settings, google_credentials_file=tmp_path / "credentials.json")

    result = build_queue_source(configured, "google", None, None)

    assert result is sentinel
    assert captured["kwargs"] == {
        "plan_worksheet": configured.google_plan_worksheet,
        "timezone_guard_enabled": True,
        "local_call_start": configured.local_call_start,
        "local_lead_cutoff": configured.local_lead_cutoff,
    }
